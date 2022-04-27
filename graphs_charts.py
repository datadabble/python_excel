import openpyxl
from openpyxl.chart import PieChart, Reference, Series, PieChart3D
wb = openpyxl.Workbook()
ws = wb.active

#create data in ws
data = [
    ['Flavor', 'Sold'],
    ['vanilla', 1500],
    ['chocolate', 1700],
    ['strawberry', 600],
    ['pumpkin', 800]
]

for rows in data:
    ws.append(rows)

chart = PieChart()
labels = Reference(ws, min_col=1, min_row=2, max_row=5)
data = Reference(ws, min_col=2, min_row=1, max_row=5)
chart.add_data(data, titles_from_data=True)
chart.set_categories(labels)
chart.title = 'Ice Cream by Flavor'

ws.add_chart(chart, 'C1')
wb.save('Pie.xlsx')

#tables
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
#from openpyxl import load_workbook

tab = Table(displayName='Table1', ref= 'A1:B5')
style = TableStyleInfo(name= 'TableStyleMed9',showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True)

tab.tableStyleInfo = style
ws.add_table(tab)
wb.save('table.xlsx')
